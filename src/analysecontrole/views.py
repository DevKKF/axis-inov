# Create your views here.
import datetime
from pprint import pprint

from openpyxl.styles import Alignment
from collections import defaultdict
from io import BytesIO
import base64
from django.contrib import admin
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.urls import reverse
from django.views.generic import TemplateView
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side
from datetime import datetime, timezone
from django.db.models import Q, Max, Count
from openpyxl.styles import PatternFill
from datetime import timedelta
import tempfile
import os
from openpyxl.utils import get_column_letter
from django.db.models import Subquery, OuterRef

from configurations.helper_config import verify_sql_query
from configurations.models import ActionLog, Secteur, Bureau, BusinessUnit, Branche, Banque, Apporteur, Devise, User, AuthGroup,  Tarif, \
    Compagnie, Garantie, GarantieFormule, Groupe, ModeReglement
from inov import settings
from production.models import Client, Aliment, Mouvement, Produit, PoliceAssureur, Police, HistoriquePolice, MouvementPolice
from analysecontrole.models import AnalysePortefeuille, ControleCommission
from production.templatetags.my_filters import money_field, convertir_date_multiformat
from shared.enum import PasswordType, Statut, StatutValidite, BaseCalculTM, StatutPaiementSinistre, TypePortefeuille

from production.templatetags.my_filters import money_field, convertir_date_multiformat, supprimer_espaces, convertir_date_jj_mm_aaaa, format_montant, money_format_mille


class AnalysePortefeuilleView(PermissionRequiredMixin,TemplateView):
    template_name = 'analyse/analyse.html'
    permission_required = "analysecontrole.view_analyseportefeuille"
    model = AnalysePortefeuille

    def get(self, request, *args, **kwargs):
        context_original = self.get_context_data(**kwargs)

        analyseportefeuille = AnalysePortefeuille.objects.all().order_by('-id')

        today = datetime.now(tz=timezone.utc)
        compagnies = Compagnie.objects.order_by('nom')
        businessunit = BusinessUnit.objects.order_by('libelle')
        business_units = BusinessUnit.objects.all().order_by('libelle')
        mode_renouvellement = ModeRenouvellement

        commercials = []
        utilisateur = User.objects.all().order_by('-first_name').exclude(is_admin_group=1)
        for user in utilisateur:
            if user.is_commercial:
                commercials.append(user)

        context_perso = {'analyseportefeuille': analyseportefeuille, 'compagnies': compagnies, 'businessunit': businessunit, 'commercials': commercials, 'business_units': business_units, 'today': today, 'mode_renouvellement': mode_renouvellement}

        context = {**context_original, **context_perso}

        return self.render_to_response(context)

    def post(self):
        pass

    def get_context_data(self, **kwargs):
        pprint(kwargs)
        return {
            **super().get_context_data(**kwargs),
            **admin.site.each_context(self.request),
            "opts": self.model._meta,
        }


# Portefeuille par compagnie
def generate_excel_portefeuille_compagnie(polices_qs, compagnie_nom=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portefeuille"

    headers = [
        "POLICE", "COMPAGNIE", "CLIENT", "TYPE DE CLIENT", "BRANCHE", "PRODUIT",
        "RECONDUCTION", "STATUT", "DATE DE RENOUVELEMENT", "DATE DE FIN",
        "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N", "COM ENCAISSEE", "COM ATTENDUE"
    ]
    sheet.append(headers)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = 28 if header not in [
            "DATE DE RENOUVELLEMENT", "DATE DE FIN", "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N"
        ] else 23

    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.auto_filter.ref = sheet.dimensions

    row_index = 2
    recap_data = {}

    def ajouter_police_dans_excel(police, titre):
        nonlocal row_index

        if titre not in recap_data:
            recap_data[titre] = {"count": 0, "ht": 0, "ttc": 0}

        dernier_historique = HistoriquePolice.objects.filter(police=police).order_by('-date_du_jour').first()

        prime_ht_n = prime_ht = prime_ttc = police_com_att = 0
        if dernier_historique:
            annee_actuelle = dernier_historique.date_du_jour.year
            historique_annee_precedente = HistoriquePolice.objects.filter(
                police=police, date_du_jour__year__lt=annee_actuelle
            ).order_by('-date_du_jour').first()
            if historique_annee_precedente:
                prime_ht_n = historique_annee_precedente.prime_ht or 0
            prime_ht = dernier_historique.prime_ht or 0
            prime_ttc = dernier_historique.prime_ttc or 0
            police_com_att = dernier_historique.commission_courtage or 0

        date_for_calcul = datetime.today().date()
        today = datetime.today().date()

        date_echeance_police = police.date_fin_effet or police.date_fin_police

        if police.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
            if date_echeance_police and date_echeance_police > today:
                difference_jours = (date_echeance_police - date_for_calcul).days
                statut = "A renouveler" if difference_jours <= 90 else police.etat_police
            else:
                statut = "NON renouvelé"
        else:
            statut = police.etat_police

        police_com_enc = sum(
            reglement.montant_com_courtage
            for quittance in Quittance.objects.filter(police=police)
            for reglement in Reglement.objects.filter(quittance=quittance, statut_commission="ENCAISSEE")
        )

        date_fin_effet = dernier_historique.date_fin_effet.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Tacite Reconduction" \
            and dernier_historique.date_fin_effet else ''

        date_fin_police = dernier_historique.date_fin_police.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Sans Tacite Reconduction" \
            and dernier_historique.date_fin_police else ''

        data = [
            police.numero,
            titre,
            police.client.nom if police.client else '',
            police.client.type_personne.libelle if police.client else '',
            police.produit.branche.nom if police.produit and police.produit.branche else '',
            police.produit.nom if police.produit else '',
            dernier_historique.mode_renouvellement if dernier_historique else '',
            statut,
            date_fin_effet,
            date_fin_police,
            prime_ht_n,
            prime_ht,
            prime_ttc,
            police_com_enc,
            police_com_att,
        ]

        # Définir la couleur de fond selon le statut
        fill_color = None
        if statut == "A renouveler":
            fill_color = PatternFill(start_color="f3ab04", end_color="f3ab04", fill_type="solid")
        elif statut == "NON renouvelé":
            fill_color = PatternFill(start_color="a39080", end_color="a39080", fill_type="solid")
        elif statut not in ["Annulé", "Résilié", "Suspendu"]:
            if date_fin_effet or date_fin_police:
                fill_color = PatternFill(start_color="dae8ec", end_color="dae8ec", fill_type="solid")
            else:
                fill_color = PatternFill(start_color="fafa35", end_color="fafa35", fill_type="solid")

        sheet.append(data)

        for col_idx, val in enumerate(data, start=1):
            cell = sheet.cell(row=row_index, column=col_idx)
            cell.border = thin_border
            if col_idx in [11, 12, 13, 14, 15]:
                cell.number_format = '#,##0'
            if col_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="center")
            if fill_color:
                cell.fill = fill_color

        recap_data[titre]["count"] += 1
        recap_data[titre]["ht"] += prime_ht
        recap_data[titre]["ttc"] += prime_ttc

        row_index += 1

    for police in polices_qs:
        compagnie_nom_police = police.compagnie.nom if police.compagnie else "Aucune compagnie"
        ajouter_police_dans_excel(police, compagnie_nom_police)

    recap_start = row_index + 3

    sheet.merge_cells(start_row=recap_start, start_column=2, end_row=recap_start, end_column=5)
    title_cell = sheet.cell(row=recap_start, column=2)
    title_cell.value = "RÉCAPITULATIF PAR COMPAGNIE"
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center")

    recap_header_row = recap_start + 1
    recap_headers = ["", "COMPAGNIE", "NOMBRE DE POLICE", "MONTANT TOTAL HT EX N", "MONTANT TTC EX N"]
    sheet.append(recap_headers)

    for col in range(2, 6):
        cell = sheet.cell(row=recap_header_row, column=col)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    total_polices = total_ht = total_ttc = 0
    recap_data_row = recap_header_row

    for compagnie_, vals in recap_data.items():
        recap_data_row += 1
        sheet.cell(row=recap_data_row, column=2, value=compagnie_)
        sheet.cell(row=recap_data_row, column=3, value=vals['count'])
        sheet.cell(row=recap_data_row, column=4, value=vals['ht']).number_format = '#,##0'
        sheet.cell(row=recap_data_row, column=5, value=vals['ttc']).number_format = '#,##0'

        for i in [4, 5]:
            sheet.cell(row=recap_data_row, column=i).number_format = '#,##0'
        for col in range(2, 6):
            sheet.cell(row=recap_data_row, column=col).border = thin_border

        total_polices += vals['count']
        total_ht += vals['ht']
        total_ttc += vals['ttc']

    total_row = recap_data_row + 2
    sheet.cell(row=total_row, column=2, value="TOTAL GENERAL").font = Font(bold=True)

    cell_polices = sheet.cell(row=total_row, column=3, value=total_polices)
    cell_polices.font = Font(bold=True)

    cell_ht = sheet.cell(row=total_row, column=4, value=total_ht)
    cell_ht.font = Font(bold=True)
    cell_ht.number_format = '#,##0'

    cell_ttc = sheet.cell(row=total_row, column=5, value=total_ttc)
    cell_ttc.font = Font(bold=True)
    cell_ttc.number_format = '#,##0'

    for col in range(2, 6):
        sheet.cell(row=total_row, column=col).border = thin_border

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def add_portefeuille_compagnie(request):
    compagnie_id = request.POST.get('search_compagnie')
    mode_renouvellement = request.POST.get("search_mode_renouvellement")
    date_requete = request.POST.get('date_requete') or datetime.today().strftime("%d/%m/%Y")

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    polices_qs = Police.objects.filter(
        client__isnull=False,
        historique_polices__isnull=False,
    ).annotate(
        dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
    )

    if compagnie_id == "TOUT":
        pass  # On garde toutes les polices pour le récapitulatif dans l'Excel
    elif compagnie_id == "AUCUN":
        polices_qs = polices_qs.filter(compagnie__isnull=True)
    else:
        polices_qs = polices_qs.filter(compagnie_id=compagnie_id)

    if mode_renouvellement:
        polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=mode_renouvellement)

    if not polices_qs.exists():
        return JsonResponse({
            'statut': 0,
            'message': "Aucune police trouvée avec ces critères."
        })

    workbook = generate_excel_portefeuille_compagnie(polices_qs.distinct())

    # Créer un fichier temporaire
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        tmp_file.write(workbook.getvalue())
        tmp_file_path = tmp_file.name

    # Supprimer le fichier temporaire après l'avoir enregistré
    os.unlink(tmp_file_path)

    filename = f"{date_requete}_Portefeuille_compagnie.xlsx"
    if compagnie_id and compagnie_id != "TOUT":
        compagnie = Compagnie.objects.filter(id=compagnie_id).first()
        if compagnie:
            filename = f"{date_requete}_Portefeuille_{compagnie.nom}.xlsx"
        elif compagnie_id == "AUCUN":
            filename = f"{date_requete}_Portefeuille_Sans_compagnie.xlsx"
    elif mode_renouvellement:
        filename = f"{date_requete}_Portefeuille_Renouvellement_{mode_renouvellement}.xlsx"
    elif compagnie_id == "TOUT":
        filename = f"{date_requete}_Portefeuille_Global.xlsx"


    return JsonResponse({
        'statut': 1,
        'message': "Portefeuille compagnie généré avec succès !",
        'data': {
            'filename': filename,
            'file_base64': base64.b64encode(workbook.getvalue()).decode()
        }
    })


# Chargement des polices liées à la compagnie
def get_client_by_compagnie(request):
    compagnie_id = request.GET.get('compagnie_id')
    search_mode_renouvellement = request.GET.get("mode_renouvellement")
    today = datetime.today().date()
    date_for_calcul = datetime.today().date()
    polices_par_compagnie = {}

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    if compagnie_id == "TOUT":
        compagnies = Compagnie.objects.all().order_by('nom')

        for compagnie in compagnies:

            polices_qs = Police.objects.filter(
                client__isnull=False,
                historique_polices__isnull=False,
                compagnie_id=compagnie.id
            ).annotate(
                dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
            ).distinct()

            if search_mode_renouvellement:
                polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

            polices = []
            for plc in polices_qs:
                dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
                dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

                etat_police = None

                date_echeance_police = None
                if plc.date_fin_effet:
                    date_echeance_police = plc.date_fin_effet
                elif plc.date_fin_police:
                    date_echeance_police = plc.date_fin_police

                if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                    if date_echeance_police and date_echeance_police > today:
                        difference_jours = (date_echeance_police - date_for_calcul).days
                        if difference_jours <= 90:
                            etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                        else:
                            etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                    else:
                        etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

                detail_url = reverse('police.details', args=[plc.id])
                numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

                polices.append({
                    'id': plc.id,
                    'nom': plc.client.nom if plc.client else '',
                    'prenoms': plc.client.prenoms if plc.client else '',
                    'numero': numero_html,
                    'produit': plc.produit.nom,
                    'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                    'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                    'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                    'statut': etat_police,
                    'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                    'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                    'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
                })

            if polices:
                polices_par_compagnie[compagnie.nom] = {
                    "polices": polices,
                }

        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            compagnie_id__isnull=True
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        autres_polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            autres_polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_compagnie["Aucune compagnie"] = {
            "polices": autres_polices,
        }

    elif compagnie_id == "AUCUN":
        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            compagnie_id__isnull=True
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_compagnie["Aucune compagnie"] = {
            "polices": polices,
        }

    else:
        compagnie = Compagnie.objects.filter(id=compagnie_id).first()

        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            compagnie_id=compagnie_id
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        if polices:
            polices_par_compagnie[compagnie.nom] = {
                "polices": polices,
            }

    return JsonResponse({
        'polices_par_compagnie': polices_par_compagnie,
    })


# Portefeuille par commercial
def generate_excel_portefeuille_commercial(polices_qs, commercial_nom=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portefeuille"

    headers = [
        "POLICE", "COMMERCIAL", "CLIENT", "TYPE DE CLIENT", "BRANCHE", "PRODUIT",
        "RECONDUCTION", "STATUT", "DATE DE RENOUVELEMENT", "DATE DE FIN",
        "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N", "COM ENCAISSEE", "COM ATTENDUE"
    ]
    sheet.append(headers)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = 28 if header not in [
            "DATE DE RENOUVELLEMENT", "DATE DE FIN", "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N"
        ] else 23

    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.auto_filter.ref = sheet.dimensions

    row_index = 2
    recap_data = {}

    def ajouter_police_dans_excel(police, titre):
        nonlocal row_index

        if titre not in recap_data:
            recap_data[titre] = {"count": 0, "ht": 0, "ttc": 0}

        dernier_historique = HistoriquePolice.objects.filter(police=police).order_by('-date_du_jour').first()

        prime_ht_n = prime_ht = prime_ttc = police_com_att = 0
        if dernier_historique:
            annee_actuelle = dernier_historique.date_du_jour.year
            historique_annee_precedente = HistoriquePolice.objects.filter(
                police=police, date_du_jour__year__lt=annee_actuelle
            ).order_by('-date_du_jour').first()
            if historique_annee_precedente:
                prime_ht_n = historique_annee_precedente.prime_ht or 0
            prime_ht = dernier_historique.prime_ht or 0
            prime_ttc = dernier_historique.prime_ttc or 0
            police_com_att = dernier_historique.commission_courtage or 0

        date_for_calcul = datetime.today().date()
        today = datetime.today().date()

        date_echeance_police = police.date_fin_effet or police.date_fin_police

        if police.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
            if date_echeance_police and date_echeance_police > today:
                difference_jours = (date_echeance_police - date_for_calcul).days
                statut = "A renouveler" if difference_jours <= 90 else police.etat_police
            else:
                statut = "NON renouvelé"
        else:
            statut = police.etat_police

        police_com_enc = sum(
            reglement.montant_com_courtage
            for quittance in Quittance.objects.filter(police=police)
            for reglement in Reglement.objects.filter(quittance=quittance, statut_commission="ENCAISSEE")
        )

        date_fin_effet = dernier_historique.date_fin_effet.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Tacite Reconduction" \
            and dernier_historique.date_fin_effet else ''

        date_fin_police = dernier_historique.date_fin_police.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Sans Tacite Reconduction" \
            and dernier_historique.date_fin_police else ''

        data = [
            police.numero,
            titre,
            police.client.nom if police.client else '',
            police.client.type_personne.libelle if police.client else '',
            police.produit.branche.nom if police.produit and police.produit.branche else '',
            police.produit.nom if police.produit else '',
            dernier_historique.mode_renouvellement if dernier_historique else '',
            statut,
            date_fin_effet,
            date_fin_police,
            prime_ht_n,
            prime_ht,
            prime_ttc,
            police_com_enc,
            police_com_att,
        ]

        # Définir la couleur de fond selon le statut
        fill_color = None
        if statut == "A renouveler":
            fill_color = PatternFill(start_color="f3ab04", end_color="f3ab04", fill_type="solid")
        elif statut == "NON renouvelé":
            fill_color = PatternFill(start_color="a39080", end_color="a39080", fill_type="solid")
        elif statut not in ["Annulé", "Résilié", "Suspendu"]:
            if date_fin_effet or date_fin_police:
                fill_color = PatternFill(start_color="dae8ec", end_color="dae8ec", fill_type="solid")
            else:
                fill_color = PatternFill(start_color="fafa35", end_color="fafa35", fill_type="solid")

        sheet.append(data)

        for col_idx, val in enumerate(data, start=1):
            cell = sheet.cell(row=row_index, column=col_idx)
            cell.border = thin_border
            if col_idx in [11, 12, 13, 14, 15]:
                cell.number_format = '#,##0'
            if col_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="center")
            if fill_color:
                cell.fill = fill_color

        recap_data[titre]["count"] += 1
        recap_data[titre]["ht"] += prime_ht
        recap_data[titre]["ttc"] += prime_ttc

        row_index += 1

    for police in polices_qs:
        commercial_nom_police = police.commercial.first_name + " " + police.commercial.last_name if police.commercial else "Aucun commercial"
        ajouter_police_dans_excel(police, commercial_nom_police)

    recap_start = row_index + 3

    sheet.merge_cells(start_row=recap_start, start_column=2, end_row=recap_start, end_column=5)
    title_cell = sheet.cell(row=recap_start, column=2)
    title_cell.value = "RÉCAPITULATIF PAR COMMERCIAL"
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center")

    recap_header_row = recap_start + 1
    recap_headers = ["", "COMMERCIAL", "NOMBRE DE POLICE", "MONTANT TOTAL HT EX N", "MONTANT TTC EX N"]
    sheet.append(recap_headers)

    for col in range(2, 6):
        cell = sheet.cell(row=recap_header_row, column=col)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    total_polices = total_ht = total_ttc = 0
    recap_data_row = recap_header_row

    for commercial, vals in recap_data.items():
        recap_data_row += 1
        sheet.cell(row=recap_data_row, column=2, value=commercial)
        sheet.cell(row=recap_data_row, column=3, value=vals['count'])
        sheet.cell(row=recap_data_row, column=4, value=vals['ht']).number_format = '#,##0'
        sheet.cell(row=recap_data_row, column=5, value=vals['ttc']).number_format = '#,##0'

        for i in [4, 5]:
            sheet.cell(row=recap_data_row, column=i).number_format = '#,##0'
        for col in range(2, 6):
            sheet.cell(row=recap_data_row, column=col).border = thin_border

        total_polices += vals['count']
        total_ht += vals['ht']
        total_ttc += vals['ttc']

    total_row = recap_data_row + 2
    sheet.cell(row=total_row, column=2, value="TOTAL GENERAL").font = Font(bold=True)

    cell_polices = sheet.cell(row=total_row, column=3, value=total_polices)
    cell_polices.font = Font(bold=True)

    cell_ht = sheet.cell(row=total_row, column=4, value=total_ht)
    cell_ht.font = Font(bold=True)
    cell_ht.number_format = '#,##0'

    cell_ttc = sheet.cell(row=total_row, column=5, value=total_ttc)
    cell_ttc.font = Font(bold=True)
    cell_ttc.number_format = '#,##0'

    for col in range(2, 6):
        sheet.cell(row=total_row, column=col).border = thin_border

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def add_portefeuille_commercial(request):
    commercial_id = request.POST.get('search_commercial')
    mode_renouvellement = request.POST.get("search_mode_renouvellement")
    date_requete = request.POST.get('date_requete') or datetime.today().strftime("%d/%m/%Y")

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    polices_qs = Police.objects.filter(
        client__isnull=False,
        historique_polices__isnull=False,
    ).annotate(
        dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
    )

    if commercial_id == "TOUT":
        pass  # On garde toutes les polices pour le récapitulatif dans l'Excel
    elif commercial_id == "AUCUN":
        polices_qs = polices_qs.filter(commercial__isnull=True)
    else:
        polices_qs = polices_qs.filter(commercial_id=commercial_id)

    if mode_renouvellement:
        polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=mode_renouvellement)

    if not polices_qs.exists():
        return JsonResponse({
            'statut': 0,
            'message': "Aucune police trouvée avec ces critères."
        })

    workbook = generate_excel_portefeuille_commercial(polices_qs.distinct())

    # Créer un fichier temporaire
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        tmp_file.write(workbook.getvalue())
        tmp_file_path = tmp_file.name

    # Supprimer le fichier temporaire après l'avoir enregistré
    os.unlink(tmp_file_path)

    filename = f"{date_requete}_Portefeuille_Commercial.xlsx"
    if commercial_id and commercial_id != "TOUT":
        commercial = User.objects.filter(id=commercial_id).first()
        if commercial:
            filename = f"{date_requete}_Portefeuille_{commercial.first_name}_{commercial.last_name}.xlsx"
        elif commercial_id == "AUCUN":
            filename = f"{date_requete}_Portefeuille_Sans_Commercial.xlsx"
    elif mode_renouvellement:
        filename = f"{date_requete}_Portefeuille_Renouvellement_{mode_renouvellement}.xlsx"
    elif commercial_id == "TOUT":
        filename = f"{date_requete}_Portefeuille_Global.xlsx"


    return JsonResponse({
        'statut': 1,
        'message': "Portefeuille commercial généré avec succès !",
        'data': {
            'filename': filename,
            'file_base64': base64.b64encode(workbook.getvalue()).decode()
        }
    })


# Chargement des polices liées au commercial
def get_client_by_commercial(request):
    commercial_id = request.GET.get('commercial_id')
    search_mode_renouvellement = request.GET.get("mode_renouvellement")
    today = datetime.today().date()
    date_for_calcul = datetime.today().date()
    polices_par_commercial = {}

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    if commercial_id == "TOUT":
        commercials = User.objects.all().order_by('first_name')

        for commercial in commercials:

            polices_qs = Police.objects.filter(
                client__isnull=False,
                historique_polices__isnull=False,
                commercial_id=commercial.id
            ).annotate(
                dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
            ).distinct()

            if search_mode_renouvellement:
                polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

            polices = []
            for plc in polices_qs:
                dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
                dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

                etat_police = None

                date_echeance_police = None
                if plc.date_fin_effet:
                    date_echeance_police = plc.date_fin_effet
                elif plc.date_fin_police:
                    date_echeance_police = plc.date_fin_police

                if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                    if date_echeance_police and date_echeance_police > today:
                        difference_jours = (date_echeance_police - date_for_calcul).days
                        if difference_jours <= 90:
                            etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                        else:
                            etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                    else:
                        etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

                detail_url = reverse('police.details', args=[plc.id])
                numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

                polices.append({
                    'id': plc.id,
                    'nom': plc.client.nom if plc.client else '',
                    'prenoms': plc.client.prenoms if plc.client else '',
                    'numero': numero_html,
                    'produit': plc.produit.nom,
                    'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                    'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                    'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                    'statut': etat_police,
                    'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                    'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                    'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
                })

            if polices:
                polices_par_commercial[commercial.first_name + ' ' + commercial.last_name] = {
                    "polices": polices,
                }

        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            commercial_id__isnull=True
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        autres_polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            autres_polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_commercial["Aucun commercial"] = {
            "polices": autres_polices,
        }

    elif commercial_id == "AUCUN":
        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            commercial_id__isnull=True
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (
                    plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime(
                    "%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(
                    dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_commercial["Aucun commercial"] = {
            "polices": polices,
        }

    else:
        commercial = User.objects.filter(id=commercial_id).first()

        polices_qs = Police.objects.filter(
            client__isnull=False,
            historique_polices__isnull=False,
            commercial_id=commercial_id
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (
                    plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime(
                    "%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(
                    dernier_historique.commission_courtage) if dernier_historique else '',
            })

        if polices:
            polices_par_commercial[commercial.first_name + ' ' + commercial.last_name] = {
                "polices": polices,
            }

    return JsonResponse({
        'polices_par_commercial': polices_par_commercial,
    })


# Portefeuille par business unit
def generate_excel_portefeuille_business_unit(polices_qs, business_unit_nom=""):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portefeuille"

    headers = [
        "POLICE", "BUSINESS UNIT", "CLIENT", "TYPE DE CLIENT", "BRANCHE", "PRODUIT",
        "RECONDUCTION", "STATUT", "DATE DE RENOUVELEMENT", "DATE DE FIN",
        "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N", "COM ENCAISSEE", "COM ATTENDUE"
    ]
    sheet.append(headers)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = 28 if header not in [
            "DATE DE RENOUVELLEMENT", "DATE DE FIN", "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N"
        ] else 23

    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.auto_filter.ref = sheet.dimensions

    row_index = 2
    recap_data = {}

    def ajouter_police_dans_excel(police, titre):
        nonlocal row_index

        if titre not in recap_data:
            recap_data[titre] = {"count": 0, "ht": 0, "ttc": 0}

        dernier_historique = HistoriquePolice.objects.filter(police=police).order_by('-date_du_jour').first()

        prime_ht_n = prime_ht = prime_ttc = police_com_att = 0
        if dernier_historique:
            annee_actuelle = dernier_historique.date_du_jour.year
            historique_annee_precedente = HistoriquePolice.objects.filter(
                police=police, date_du_jour__year__lt=annee_actuelle
            ).order_by('-date_du_jour').first()
            if historique_annee_precedente:
                prime_ht_n = historique_annee_precedente.prime_ht or 0
            prime_ht = dernier_historique.prime_ht or 0
            prime_ttc = dernier_historique.prime_ttc or 0
            police_com_att = dernier_historique.commission_courtage or 0

        date_for_calcul = datetime.today().date()
        today = datetime.today().date()

        date_echeance_police = police.date_fin_effet or police.date_fin_police

        if police.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
            if date_echeance_police and date_echeance_police > today:
                difference_jours = (date_echeance_police - date_for_calcul).days
                statut = "A renouveler" if difference_jours <= 90 else police.etat_police
            else:
                statut = "NON renouvelé"
        else:
            statut = police.etat_police

        police_com_enc = sum(
            reglement.montant_com_courtage
            for quittance in Quittance.objects.filter(police=police)
            for reglement in Reglement.objects.filter(quittance=quittance, statut_commission="ENCAISSEE")
        )

        date_fin_effet = dernier_historique.date_fin_effet.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Tacite Reconduction" \
            and dernier_historique.date_fin_effet else ''

        date_fin_police = dernier_historique.date_fin_police.strftime('%d/%m/%Y') \
            if dernier_historique and dernier_historique.mode_renouvellement == "Sans Tacite Reconduction" \
            and dernier_historique.date_fin_police else ''

        data = [
            police.numero,
            titre,
            police.client.nom if police.client else '',
            police.client.type_personne.libelle if police.client else '',
            police.produit.branche.nom if police.produit and police.produit.branche else '',
            police.produit.nom if police.produit else '',
            dernier_historique.mode_renouvellement if dernier_historique else '',
            statut,
            date_fin_effet,
            date_fin_police,
            prime_ht_n,
            prime_ht,
            prime_ttc,
            police_com_enc,
            police_com_att,
        ]

        # Définir la couleur de fond selon le statut
        fill_color = None
        if statut == "A renouveler":
            fill_color = PatternFill(start_color="f3ab04", end_color="f3ab04", fill_type="solid")
        elif statut == "NON renouvelé":
            fill_color = PatternFill(start_color="a39080", end_color="a39080", fill_type="solid")
        elif statut not in ["Annulé", "Résilié", "Suspendu"]:
            if date_fin_effet or date_fin_police:
                fill_color = PatternFill(start_color="dae8ec", end_color="dae8ec", fill_type="solid")
            else:
                fill_color = PatternFill(start_color="fafa35", end_color="fafa35", fill_type="solid")

        sheet.append(data)

        for col_idx, val in enumerate(data, start=1):
            cell = sheet.cell(row=row_index, column=col_idx)
            cell.border = thin_border
            if col_idx in [11, 12, 13, 14, 15]:
                cell.number_format = '#,##0'
            if col_idx in [9, 10]:
                cell.alignment = Alignment(horizontal="center")
            if fill_color:
                cell.fill = fill_color

        recap_data[titre]["count"] += 1
        recap_data[titre]["ht"] += prime_ht
        recap_data[titre]["ttc"] += prime_ttc

        row_index += 1

    for police in polices_qs:
        business_unit_nom_police = police.client.business_unit.libelle if police.client.business_unit else "Aucun business unit"
        ajouter_police_dans_excel(police, business_unit_nom_police)

    recap_start = row_index + 3

    sheet.merge_cells(start_row=recap_start, start_column=2, end_row=recap_start, end_column=5)
    title_cell = sheet.cell(row=recap_start, column=2)
    title_cell.value = "RÉCAPITULATIF PAR BUSINESS UNIT"
    title_cell.font = Font(bold=True, size=12)
    title_cell.alignment = Alignment(horizontal="center")

    recap_header_row = recap_start + 1
    recap_headers = ["", "BUSINESS UNIT", "NOMBRE DE POLICE", "MONTANT TOTAL HT EX N", "MONTANT TTC EX N"]
    sheet.append(recap_headers)

    for col in range(2, 6):
        cell = sheet.cell(row=recap_header_row, column=col)
        cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    total_polices = total_ht = total_ttc = 0
    recap_data_row = recap_header_row

    for business_unit_, vals in recap_data.items():
        recap_data_row += 1
        sheet.cell(row=recap_data_row, column=2, value=business_unit_)
        sheet.cell(row=recap_data_row, column=3, value=vals['count'])
        sheet.cell(row=recap_data_row, column=4, value=vals['ht']).number_format = '#,##0'
        sheet.cell(row=recap_data_row, column=5, value=vals['ttc']).number_format = '#,##0'

        for i in [4, 5]:
            sheet.cell(row=recap_data_row, column=i).number_format = '#,##0'
        for col in range(2, 6):
            sheet.cell(row=recap_data_row, column=col).border = thin_border

        total_polices += vals['count']
        total_ht += vals['ht']
        total_ttc += vals['ttc']

    total_row = recap_data_row + 2
    sheet.cell(row=total_row, column=2, value="TOTAL GENERAL").font = Font(bold=True)

    cell_polices = sheet.cell(row=total_row, column=3, value=total_polices)
    cell_polices.font = Font(bold=True)

    cell_ht = sheet.cell(row=total_row, column=4, value=total_ht)
    cell_ht.font = Font(bold=True)
    cell_ht.number_format = '#,##0'

    cell_ttc = sheet.cell(row=total_row, column=5, value=total_ttc)
    cell_ttc.font = Font(bold=True)
    cell_ttc.number_format = '#,##0'

    for col in range(2, 6):
        sheet.cell(row=total_row, column=col).border = thin_border

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output



def add_portefeuille_business_unit(request):
    business_unit_id = request.POST.get('search_business_unit')
    mode_renouvellement = request.POST.get("search_mode_renouvellement")
    date_requete = request.POST.get('date_requete') or datetime.today().strftime("%d/%m/%Y")

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    polices_qs = Police.objects.filter(
        client__isnull=False,
        historique_polices__isnull=False,
    ).annotate(
        dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
    )

    if business_unit_id == "TOUT":
        pass  # On garde toutes les polices pour le récapitulatif dans l'Excel
    elif business_unit_id == "AUCUN":
        polices_qs = polices_qs.filter(client__business_unit__isnull=True)
    else:
        polices_qs = polices_qs.filter(client__business_unit_id=business_unit_id)

    if mode_renouvellement:
        polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=mode_renouvellement)

    if not polices_qs.exists():
        return JsonResponse({
            'statut': 0,
            'message': "Aucune police trouvée avec ces critères."
        })

    workbook = generate_excel_portefeuille_business_unit(polices_qs.distinct())

    # Créer un fichier temporaire
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        tmp_file.write(workbook.getvalue())
        tmp_file_path = tmp_file.name

    # Supprimer le fichier temporaire après l'avoir enregistré
    os.unlink(tmp_file_path)

    filename = f"{date_requete}_Portefeuille_BusinessUnit.xlsx"
    if business_unit_id and business_unit_id != "TOUT":
        business_unit = BusinessUnit.objects.filter(id=business_unit_id).first()
        if business_unit:
            filename = f"{date_requete}_Portefeuille_{business_unit.libelle}.xlsx"
        elif business_unit_id == "AUCUN":
            filename = f"{date_requete}_Portefeuille_Sans_BusinessUnit.xlsx"
    elif mode_renouvellement:
        filename = f"{date_requete}_Portefeuille_Renouvellement_{mode_renouvellement}.xlsx"
    elif business_unit_id == "TOUT":
        filename = f"{date_requete}_Portefeuille_Global.xlsx"


    return JsonResponse({
        'statut': 1,
        'message': "Portefeuille business_unit généré avec succès !",
        'data': {
            'filename': filename,
            'file_base64': base64.b64encode(workbook.getvalue()).decode()
        }
    })


# Chargement des polices liées au business unit
def get_client_by_business_unit(request):
    business_unit_id = request.GET.get('business_unit_id')
    search_mode_renouvellement = request.GET.get("mode_renouvellement")
    today = datetime.today().date()
    date_for_calcul = datetime.today().date()
    polices_par_business_unit = {}

    dernier_historique_subquery = HistoriquePolice.objects.filter(
        police_id=OuterRef('pk')
    ).order_by('-date_du_jour').values('mode_renouvellement')[:1]

    if business_unit_id == "TOUT":
        business_units = BusinessUnit.objects.all().order_by('libelle')

        for business_unit in business_units:

            polices_qs = Police.objects.filter(
                client__business_unit_id=business_unit.id,
                historique_polices__isnull=False
            ).annotate(
                dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
            ).distinct()

            if search_mode_renouvellement:
                polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

            polices = []
            for plc in polices_qs:
                dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
                dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

                etat_police = None

                date_echeance_police = None
                if plc.date_fin_effet:
                    date_echeance_police = plc.date_fin_effet
                elif plc.date_fin_police:
                    date_echeance_police = plc.date_fin_police

                if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                    if date_echeance_police and date_echeance_police > today:
                        difference_jours = (date_echeance_police - date_for_calcul).days
                        if difference_jours <= 90:
                            etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                        else:
                            etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                    else:
                        etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

                detail_url = reverse('police.details', args=[plc.id])
                numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

                polices.append({
                    'id': plc.id,
                    'nom': plc.client.nom if plc.client else '',
                    'prenoms': plc.client.prenoms if plc.client else '',
                    'numero': numero_html,
                    'produit': plc.produit.nom,
                    'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                    'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                    'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                    'statut': etat_police,
                    'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                    'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                    'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
                })

            if polices:
                polices_par_business_unit[business_unit.libelle] = {
                    "polices": polices,
                }

        polices_qs = Police.objects.filter(
            client__business_unit_id__isnull=True,
            historique_polices__isnull=False
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        autres_polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            autres_polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_business_unit["Aucun business unit"] = {
            "polices": autres_polices,
        }

    elif business_unit_id == "AUCUN":
        polices_qs = Police.objects.filter(
            client__business_unit_id__isnull=True,
            historique_polices__isnull=False
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (
                    plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime(
                    "%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(
                    dernier_historique.commission_courtage) if dernier_historique else '',
            })

        polices_par_business_unit["Aucun business unit"] = {
            "polices": polices,
        }

    else:
        business_unit = BusinessUnit.objects.filter(id=business_unit_id).first()

        polices_qs = Police.objects.filter(
            client__business_unit_id=business_unit.id,
            historique_polices__isnull=False,
        ).annotate(
            dernier_mode_renouvellement=Subquery(dernier_historique_subquery)
        ).distinct()

        if search_mode_renouvellement:
            polices_qs = polices_qs.filter(dernier_mode_renouvellement__iexact=search_mode_renouvellement)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            etat_police = None

            date_echeance_police = None
            if plc.date_fin_effet:
                date_echeance_police = plc.date_fin_effet
            elif plc.date_fin_police:
                date_echeance_police = plc.date_fin_police

            if plc.etat_police not in ["Annulé", "Résilié", "Suspendu"]:
                if date_echeance_police and date_echeance_police > today:
                    difference_jours = (date_echeance_police - date_for_calcul).days
                    if difference_jours <= 90:
                        etat_police = f'<span class="badge badge-warning">A renouveler</span>'
                    else:
                        etat_police = f'<span class="badge badge-success">{plc.etat_police}</span>'
                else:
                    etat_police = f'<span class="badge badge-danger">NON renouvelé</span>'
            else:
                etat_police = f'<span class="badge badge-danger">{plc.etat_police}</span>'

            detail_url = reverse('police.details', args=[plc.id])
            numero_html = f'<a href="{detail_url}" class="text-center bouton_action" style="color:#F16623;" target="_blank">{plc.numero}</a>'

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': numero_html,
                'produit': plc.produit.nom,
                'date_debut_effet': plc.date_debut_effet.strftime("%d/%m/%Y") if plc.date_debut_effet else '',
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else (plc.date_fin_police.strftime("%d/%m/%Y") if plc.date_fin_police else None),
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'mode_renouvellement': dernier_historique.mode_renouvellement if dernier_historique else '',
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'commission_courtage': money_field(dernier_historique.commission_courtage) if dernier_historique else '',
            })

        if polices:
            polices_par_business_unit[business_unit.libelle] = {
                "polices": polices,
            }

    return JsonResponse({
        'polices_par_business_unit': polices_par_business_unit,
    })



def commande_perso(request):
    compagnies = PoliceAssureur.objects.filter(type_compagnie_id=1)

    for cie in compagnies:
        historiques = HistoriquePolice.objects.filter(id=cie.historique_police_id)

        for histo in historiques:
            # Mise à jours de la table historique_police
            histo.compagnie_id = cie.compagnie_id
            histo.updated_by = request.user
            histo.save()

            # Récupération de la police
            Police.objects.filter(id=histo.police_id).update(
                compagnie_id=cie.compagnie_id,
                updated_by = request.user
            )

    return JsonResponse({
        'statut': 1,
        'message': "Police et son historique mis à jour avec succès !",
        'data': {}
    })